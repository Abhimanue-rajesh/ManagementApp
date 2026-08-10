from datetime import date, datetime
from pathlib import Path

from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from tasks.models import (
    Project,
    ProjectHistory,
    Task,
    TaskActivity,
)


class Command(BaseCommand):
    help = "Import projects, tasks and task activities from Software Task Excel."

    STATUS_MAP = {
        "pending": "not_started",
        "planned": "not_started",
        "not started": "not_started",
        "not_started": "not_started",
        "in progress": "in_progress",
        "in_progress": "in_progress",
        "ongoing": "in_progress",
        "waiting for approval": "waiting_for_approval",
        "waiting_for_approval": "waiting_for_approval",
        "terminated": "terminated",
        "done": "closed",
        "completed": "closed",
        "complete": "closed",
        "closed": "closed",
    }

    # Column numbers from the actual uploaded workbook.
    #
    # Column 12 has an Excel header which was interpreted as
    # 2026-03-08, but from the chronological layout it appears
    # to represent the 3 Aug update.
    #
    # Column 13 contains activity data although its header is blank.
    ACTIVITY_COLUMNS = {
        6: None,  # "Updation" - no reliable date in header
        8: date(2026, 7, 20),
        9: date(2026, 7, 20),
        10: date(2026, 7, 22),
        11: date(2026, 7, 31),
        12: date(2026, 8, 3),
        13: None,
        14: date(2026, 8, 6),
    }

    def add_arguments(self, parser):
        parser.add_argument(
            "excel_file",
            type=str,
            help="Path to the Excel file.",
        )

        parser.add_argument(
            "--user",
            required=True,
            help=(
                "Username of the Django user who will own imported "
                "tasks and be recorded as project creator."
            ),
        )

        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Validate/import inside a transaction and roll everything back.",
        )

    def handle(self, *args, **options):
        excel_path = Path(options["excel_file"])
        username = options["user"]
        dry_run = options["dry_run"]

        if not excel_path.exists():
            raise CommandError(f"Excel file does not exist: {excel_path}")

        User = get_user_model()

        try:
            import_user = User.objects.get(username=username)
        except User.DoesNotExist:
            raise CommandError(f'Django user "{username}" does not exist.')

        try:
            workbook = load_workbook(
                excel_path,
                data_only=True,
            )
        except Exception as exc:
            raise CommandError(f"Unable to open Excel file: {exc}")

        if "Main" not in workbook.sheetnames:
            raise CommandError('Expected a worksheet named "Main".')

        sheet = workbook["Main"]

        stats = {
            "projects_created": 0,
            "projects_existing": 0,
            "tasks_created": 0,
            "tasks_updated": 0,
            "tasks_skipped": 0,
            "activities_created": 0,
            "activities_existing": 0,
            "rows_skipped": 0,
        }

        self.stdout.write(self.style.NOTICE(f"Reading: {excel_path}"))

        self.stdout.write(f"Worksheet: {sheet.title}")

        with transaction.atomic():
            for row_number in range(2, sheet.max_row + 1):
                module_value = sheet.cell(
                    row=row_number,
                    column=2,
                ).value

                task_value = sheet.cell(
                    row=row_number,
                    column=3,
                ).value

                owner_value = sheet.cell(
                    row=row_number,
                    column=4,
                ).value

                status_value = sheet.cell(
                    row=row_number,
                    column=5,
                ).value

                deadline_value = sheet.cell(
                    row=row_number,
                    column=7,
                ).value

                module_name = self.clean_text(module_value)
                raw_task = self.clean_text(task_value)
                owner = self.clean_text(owner_value)
                raw_status = self.clean_text(status_value)

                # Completely empty row
                if not module_name and not raw_task:
                    continue

                # Module exists but there is no task.
                #
                # Example in current spreadsheet:
                # "crm feedback form"
                if module_name and not raw_task:
                    stats["rows_skipped"] += 1

                    self.stdout.write(
                        self.style.WARNING(
                            f"Row {row_number}: skipped - "
                            f'project "{module_name}" has no task.'
                        )
                    )

                    continue

                if not module_name:
                    stats["rows_skipped"] += 1

                    self.stdout.write(
                        self.style.WARNING(
                            f"Row {row_number}: skipped - "
                            "task has no Module/Project."
                        )
                    )

                    continue

                project, project_created = self.get_or_create_project(
                    module_name=module_name,
                    import_user=import_user,
                )

                if project_created:
                    stats["projects_created"] += 1

                    ProjectHistory.objects.create(
                        project=project,
                        action="created",
                        new_status=project.status,
                        description=(
                            "Project created from Software Task Excel import."
                        ),
                        updated_by=import_user,
                    )
                else:
                    stats["projects_existing"] += 1

                title, description = self.prepare_task_text(raw_task)

                status = self.convert_status(
                    raw_status,
                    row_number,
                )

                due_date, deadline_remark = self.extract_due_date(deadline_value)

                if deadline_remark:
                    if description:
                        description = (
                            f"{description}\n\n"
                            f"Deadline / Remarks:\n"
                            f"{deadline_remark}"
                        )
                    else:
                        description = f"Deadline / Remarks:\n" f"{deadline_remark}"

                task = self.find_existing_task(
                    project=project,
                    title=title,
                )

                if task:
                    changed = False

                    # Preserve existing manually entered values
                    # unless they're currently empty/default-like.

                    if owner and not task.pending_with:
                        task.pending_with = owner
                        changed = True

                    if due_date and not task.due_date:
                        task.due_date = due_date
                        changed = True

                    if description and not task.description:
                        task.description = description
                        changed = True

                    if changed:
                        task.save()

                    stats["tasks_updated"] += 1

                else:
                    task = Task.objects.create(
                        project=project,
                        user=import_user,
                        title=title,
                        description=description,
                        status=status,
                        priority="medium",
                        pending_with=owner or None,
                        due_date=due_date,
                    )

                    stats["tasks_created"] += 1

                activity_count = self.import_activities(
                    sheet=sheet,
                    row_number=row_number,
                    task=task,
                )

                stats["activities_created"] += activity_count["created"]

                stats["activities_existing"] += activity_count["existing"]

            if dry_run:
                transaction.set_rollback(True)

                self.stdout.write("")
                self.stdout.write(
                    self.style.WARNING(
                        "DRY RUN: All database changes were rolled back."
                    )
                )

        self.print_summary(stats, dry_run)

    # ---------------------------------------------------------
    # Project
    # ---------------------------------------------------------

    def get_or_create_project(
        self,
        module_name,
        import_user,
    ):
        project = Project.objects.filter(name__iexact=module_name).first()

        if project:
            return project, False

        project = Project.objects.create(
            name=module_name,
            status="in_progress",
            created_by=import_user,
        )

        return project, True

    # ---------------------------------------------------------
    # Task
    # ---------------------------------------------------------

    def find_existing_task(
        self,
        project,
        title,
    ):
        return Task.objects.filter(
            project=project,
            title__iexact=title,
        ).first()

    def prepare_task_text(self, raw_task):
        """
        Task.title is limited to 200 characters.

        Most spreadsheet rows contain a normal task title.
        One row (POS ERP) contains an entire requirement document
        inside the Task column.

        If the cell is too long, use the first meaningful line as
        the title and store the full cell in description.
        """

        text = raw_task.strip()

        if len(text) <= 200:
            return text, ""

        lines = [line.strip() for line in text.splitlines() if line.strip()]

        if lines:
            first_line = lines[0]

            if len(first_line) <= 200:
                return first_line, text

        title = text[:197].rstrip() + "..."

        return title, text

    # ---------------------------------------------------------
    # Status
    # ---------------------------------------------------------

    def convert_status(
        self,
        raw_status,
        row_number,
    ):
        if not raw_status:
            return "not_started"

        normalized = raw_status.strip().lower().replace("\n", " ")

        if normalized in self.STATUS_MAP:
            return self.STATUS_MAP[normalized]

        # Some spreadsheet Status cells contain comments rather
        # than an actual status. Don't guess.
        self.stdout.write(
            self.style.WARNING(
                f"Row {row_number}: unknown status "
                f'"{raw_status}" -> using Not Started.'
            )
        )

        return "not_started"

    # ---------------------------------------------------------
    # Deadline
    # ---------------------------------------------------------

    def extract_due_date(self, value):
        """
        Return:
            (date or None, remaining remark or "")
        """

        if value in (None, ""):
            return None, ""

        if isinstance(value, datetime):
            return value.date(), ""

        if isinstance(value, date):
            return value, ""

        text = self.clean_text(value)

        if not text:
            return None, ""

        # Only convert clear calendar dates.
        formats = (
            "%d/%m/%Y",
            "%d/%m/%y",
            "%d-%m-%Y",
            "%d-%m-%y",
            "%Y-%m-%d",
        )

        for date_format in formats:
            try:
                parsed = datetime.strptime(
                    text,
                    date_format,
                ).date()

                return parsed, ""
            except ValueError:
                pass

        # "This week", "next week", comments, etc. are preserved
        # as description instead of inventing a date.
        return None, text

    # ---------------------------------------------------------
    # Activities
    # ---------------------------------------------------------

    def import_activities(
        self,
        sheet,
        row_number,
        task,
    ):
        created = 0
        existing = 0

        for column, fixed_date in self.ACTIVITY_COLUMNS.items():
            value = sheet.cell(
                row=row_number,
                column=column,
            ).value

            note = self.clean_text(value)

            if not note:
                continue

            activity_date = (
                fixed_date or self.detect_date_from_text(note) or timezone.localdate()
            )

            activity_exists = TaskActivity.objects.filter(
                task=task,
                activity_note=note,
                activity_date=activity_date,
            ).exists()

            if activity_exists:
                existing += 1
                continue

            TaskActivity.objects.create(
                task=task,
                activity_note=note,
                activity_date=activity_date,
            )

            created += 1

        return {
            "created": created,
            "existing": existing,
        }

    def detect_date_from_text(self, text):
        """
        Very conservative date extraction.

        We only parse a standalone obvious DD/MM/YY or DD/MM/YYYY
        date if the entire cell is the date.

        We intentionally don't guess dates hidden inside sentences.
        """

        value = text.strip()

        formats = (
            "%d/%m/%Y",
            "%d/%m/%y",
            "%d-%m-%Y",
            "%d-%m-%y",
        )

        for date_format in formats:
            try:
                return datetime.strptime(
                    value,
                    date_format,
                ).date()
            except ValueError:
                pass

        return None

    # ---------------------------------------------------------
    # Utilities
    # ---------------------------------------------------------

    def clean_text(self, value):
        if value is None:
            return ""

        if isinstance(value, str):
            return value.strip()

        return str(value).strip()

    def print_summary(
        self,
        stats,
        dry_run,
    ):
        self.stdout.write("")
        self.stdout.write(
            self.style.SUCCESS(
                "Import completed." if not dry_run else "Dry-run completed."
            )
        )

        self.stdout.write("----------------------------------------")

        self.stdout.write(f"Projects created:       " f'{stats["projects_created"]}')

        self.stdout.write(f"Projects already exist: " f'{stats["projects_existing"]}')

        self.stdout.write(f"Tasks created:          " f'{stats["tasks_created"]}')

        self.stdout.write(f"Tasks matched/existing: " f'{stats["tasks_updated"]}')

        self.stdout.write(f"Activities created:     " f'{stats["activities_created"]}')

        self.stdout.write(f"Activities existing:    " f'{stats["activities_existing"]}')

        self.stdout.write(f"Rows skipped:           " f'{stats["rows_skipped"]}')

        self.stdout.write("----------------------------------------")

        if dry_run:
            self.stdout.write(self.style.WARNING("No database records were changed."))
